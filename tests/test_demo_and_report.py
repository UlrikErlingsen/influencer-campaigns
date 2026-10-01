from io import BytesIO

from openpyxl import load_workbook

from influencesignal import DISCLAIMER
from influencesignal.compliance import STATUS_ISSUE, checklist_status
from influencesignal.demo import DEMO_NOTICE, demo_creators, load_demo
from influencesignal.report import build_html, build_report, build_xlsx
from influencesignal.storage import STAGES, Store


def _snapshot(store: Store) -> tuple:
    deliverables = [store.deliverables(int(cid)).drop(columns=["id", "engagement_id", "creator_id"]) for cid in store.campaigns()["id"]]
    return (
        store.creators().drop(columns=["id"]).to_csv(),
        store.campaigns().drop(columns=["id"]).to_csv(),
        tuple(frame.to_csv() for frame in deliverables),
    )


def test_demo_is_deterministic(tmp_path, rules) -> None:
    first, second = Store(tmp_path / "a.db"), Store(tmp_path / "b.db")
    load_demo(first, rules)
    load_demo(second, rules)
    assert _snapshot(first) == _snapshot(second)
    assert demo_creators() == demo_creators()


def test_demo_shape_and_fictional_markers(demo_store: Store) -> None:
    creators = demo_store.creators()
    assert len(creators) == 25
    assert creators["is_demo"].all()
    handles = [h for p in ("instagram", "tiktok", "youtube", "snapchat") for h in creators[p] if h]
    assert handles and all(handle.startswith("demo_") for handle in handles)
    assert all(email.endswith("@example.com") for email in creators["contact_email"])
    assert all("(demo)" in name for name in creators["name"])
    campaigns = demo_store.campaigns()
    assert len(campaigns) == 2
    assert all(".example" in url for url in campaigns["landing_url"])
    assert "Fjellbrus" in DEMO_NOTICE and "no real person" in DEMO_NOTICE


def test_demo_has_mixed_pipeline_and_a_visible_missing_label(demo_store: Store, rules) -> None:
    stages = set()
    issues = []
    for campaign_id in demo_store.campaigns()["id"].astype(int):
        campaign = demo_store.campaign(campaign_id)
        stages |= set(demo_store.engagements(campaign_id)["stage"])
        for row in demo_store.deliverables(campaign_id).to_dict("records"):
            status = checklist_status(row, campaign, demo_store.answers(int(row["id"])), rules)
            if status.label == STATUS_ISSUE:
                issues.append((row, status))
    assert stages == set(STAGES)
    assert len(issues) == 1
    row, status = issues[0]
    assert "ad_identified" in status.failing
    assert row["stage"] not in ("Paid", "Reported")


def test_every_demo_paid_creator_passed_the_gate(demo_store: Store, rules) -> None:
    for campaign_id in demo_store.campaigns()["id"].astype(int):
        engagements = demo_store.engagements(campaign_id)
        for engagement_id in engagements.loc[engagements["stage"].isin(["Paid", "Reported"]), "id"]:
            allowed, reasons = demo_store.engagement_gate(int(engagement_id), rules)
            assert allowed, reasons


def test_report_xlsx_and_html(demo_store: Store, rules) -> None:
    campaign_id = int(demo_store.campaigns().query("slug == 'fjellbrus-hostfjell-2026'")["id"].iloc[0])
    report = build_report(demo_store, campaign_id, rules)
    assert report.is_demo
    assert report.compliance_counts.get(STATUS_ISSUE) == 1

    workbook = load_workbook(BytesIO(build_xlsx(report)))
    assert {"read_me", "campaign_totals", "creators", "deliverables", "compliance_checklist", "rules_and_sources",
            "notes"} <= set(workbook.sheetnames)
    read_me = {row[0]: row[1] for row in workbook["read_me"].iter_rows(min_row=2, values_only=True)}
    assert read_me["compliance_disclaimer"] == DISCLAIMER
    assert read_me["demo_notice"] == DEMO_NOTICE
    sources = [row[4] for row in workbook["rules_and_sources"].iter_rows(min_row=2, values_only=True)]
    assert any(str(url).startswith("https://www.forbrukertilsynet.no/") for url in sources)

    html = build_html(report)
    assert DISCLAIMER in html
    assert "Fictional demo" in html
    assert "Fjellbrus Høstfjell 2026" in html
    assert "compliant" not in html.lower()
    assert "<script" not in html.lower()


def test_report_escapes_user_text(store: Store, rules) -> None:
    campaign_id = store.add_campaign({"name": "<script>alert(1)</script>", "goal": "sales"})
    html = build_html(build_report(store, campaign_id, rules))
    assert "<script>alert" not in html
    assert "&lt;script&gt;" in html


def test_report_for_empty_campaign(store: Store, rules) -> None:
    campaign_id = store.add_campaign({"name": "Empty", "goal": "awareness"})
    report = build_report(store, campaign_id, rules)
    assert build_xlsx(report)
    assert "Empty" in build_html(report)
